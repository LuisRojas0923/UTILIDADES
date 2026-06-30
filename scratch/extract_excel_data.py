import openpyxl
import os

file_path = r"c:\Users\amejoramiento6\Desktop\UTILIDADES\Cuentas entrada y salidas.xlsm"

def analyze_excel(path):
    if not os.path.exists(path):
        print(f"File not found: {path}")
        return

    try:
        # Load workbook in formula mode (data_only=False by default)
        wb = openpyxl.load_workbook(path, keep_vba=True)
        print(f"Successfully loaded: {os.path.basename(path)}")
        print(f"Sheets: {wb.sheetnames}\n")

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            print(f"--- Sheet: {sheet_name} ---")
            print(f"Dimensions: {sheet.dimensions}")
            
            # Read first 20 rows to get an idea of the structure
            rows = list(sheet.iter_rows(max_row=20))
            if not rows:
                print("Empty sheet.")
                continue

            for row in rows:
                row_data = []
                for cell in row:
                    if cell.value is not None:
                        # Check if it's a formula
                        if isinstance(cell.value, str) and cell.value.startswith('='):
                            row_data.append(f"[{cell.coordinate} FORMULA: {cell.value}]")
                        else:
                            row_data.append(f"[{cell.coordinate}: {cell.value}]")
                if row_data:
                    print(" ".join(row_data))
            print("\n")

    except Exception as e:
        print(f"An error occurred: {e}")

if __name__ == "__main__":
    analyze_excel(file_path)
